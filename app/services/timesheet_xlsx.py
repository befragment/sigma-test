"""Рендер табеля в xlsx по образцу листа «Табель» из ТЗ. Без обращения к БД: на входе готовый Timesheet."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.models import Timesheet

FONT = "Arial"
MAX_DAYS = 31
HEADER_ROW = 2
FIRST_ROW = 3
FIRST_DAY_COL = 2  # B — 1-е число
TOTAL_COL = FIRST_DAY_COL + MAX_DAYS  # AG — «кол. смен»

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D8D8D8")
_MISSING_DAY_FILL = PatternFill("solid", fgColor="A6A6A6")
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def day_column(day: int) -> str:
    return get_column_letter(FIRST_DAY_COL + day - 1)


def build_xlsx(timesheet: Timesheet) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"

    last_day_col = get_column_letter(FIRST_DAY_COL + MAX_DAYS - 1)
    total_col = get_column_letter(TOTAL_COL)
    first_data_row = FIRST_ROW
    last_data_row = FIRST_ROW + len(timesheet.rows) - 1
    total_row = last_data_row + 1

    # Строка 1: заголовок над днями (B1:AG1 объединены, как в образце).
    ws.merge_cells(start_row=1, start_column=FIRST_DAY_COL, end_row=1, end_column=TOTAL_COL)
    title = ws.cell(row=1, column=FIRST_DAY_COL, value=timesheet.title)
    title.font = Font(name=FONT, size=14, bold=True)
    title.alignment = _CENTER

    # Строка 2: шапка.
    header_font = Font(name=FONT, size=12, bold=True)
    headers = {1: "Ф.И.О.", TOTAL_COL: "кол.\nсмен"}
    headers.update({FIRST_DAY_COL + day - 1: day for day in range(1, MAX_DAYS + 1)})
    for col, value in headers.items():
        cell = ws.cell(row=HEADER_ROW, column=col, value=value)
        cell.font = header_font
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=HEADER_ROW, column=1).fill = _HEADER_FILL
    ws.cell(row=HEADER_ROW, column=TOTAL_COL).fill = _HEADER_FILL

    # Строки сотрудников: 1 в днях с отметкой, сумма по строке.
    body_font = Font(name=FONT, size=12)
    for offset, row in enumerate(timesheet.rows):
        r = first_data_row + offset
        name = ws.cell(row=r, column=1, value=row.full_name)
        name.font, name.border, name.alignment = body_font, _BORDER, _LEFT
        for day in range(1, MAX_DAYS + 1):
            cell = ws.cell(row=r, column=FIRST_DAY_COL + day - 1, value=1 if day in row.days else None)
            cell.font, cell.border, cell.alignment = body_font, _BORDER, _CENTER
        total = ws.cell(row=r, column=TOTAL_COL, value=f"=SUM({day_column(1)}{r}:{last_day_col}{r})")
        total.font, total.border, total.alignment = body_font, _BORDER, _CENTER

    # Итоговая строка: сумма по каждому столбцу дней и по «кол. смен».
    total_font = Font(name=FONT, size=12, bold=True)
    label = ws.cell(row=total_row, column=1, value="ИТОГ:")
    label.font, label.border, label.alignment = total_font, _BORDER, _LEFT
    for col in range(FIRST_DAY_COL, TOTAL_COL + 1):
        letter = get_column_letter(col)
        # Без сотрудников диапазон пуст — формула сослалась бы на саму итоговую строку.
        value = f"=SUM({letter}{first_data_row}:{letter}{last_data_row})" if timesheet.rows else 0
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.font, cell.border, cell.alignment = total_font, _BORDER, _CENTER

    # Несуществующие дни месяца (например, 30–31 февраля) — серым на всю высоту таблицы.
    for day in range(timesheet.days_in_month + 1, MAX_DAYS + 1):
        for r in range(HEADER_ROW, total_row + 1):
            ws.cell(row=r, column=FIRST_DAY_COL + day - 1).fill = _MISSING_DAY_FILL

    # Ширина в Excel меряется символами шрифта по умолчанию; Arial 12 примерно на треть шире.
    ws.column_dimensions["A"].width = max([20.57, *(len(row.full_name) * 1.3 + 2 for row in timesheet.rows)])
    for col in range(FIRST_DAY_COL, TOTAL_COL):
        ws.column_dimensions[get_column_letter(col)].width = 4.71
    ws.column_dimensions[total_col].width = 9.14
    ws.row_dimensions[HEADER_ROW].height = 30
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=FIRST_DAY_COL)

    # openpyxl не сохраняет вычисленные значения формул: Excel/LibreOffice пересчитают их при открытии
    # (флаг и так включён по умолчанию, ставим явно). Превью без пересчёта (Quick Look и т.п.) покажут 0.
    wb.calculation.fullCalcOnLoad = True

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
