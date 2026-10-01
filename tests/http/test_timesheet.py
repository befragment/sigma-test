"""Табель через HTTP на реальной БД: от сообщения до единицы в ячейке xlsx."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.domain.models import Message

CHAT_ID = -1001234567890
OTHER_CHAT_ID = -1009876543210


def sheet(response):
    return load_workbook(BytesIO(response.content))["Табель"]


async def photo(app, chat_id: int, message_id: int, caption: str, sent_at: datetime) -> None:
    await app.state.services.ingest.ingest(Message(chat_id, message_id, sent_at, caption=caption))


async def test_example_from_task_lands_in_timesheet(app, client):
    """Пример из ТЗ: «Мехоношин» в 08:50 по Москве 19.09.2025 → 1 в колонке 19."""
    await client.post("/groups", json={"chat_id": CHAT_ID, "title": "Пост 1",
                                       "window_start": "07:00", "window_end": "10:00"})
    await client.post("/employees", json={"full_name": "Мехоношин Алексей"})
    await client.post("/employees", json={"full_name": "Антонов Иван"})
    await photo(app, CHAT_ID, 10, "Мехоношин", datetime(2025, 9, 19, 5, 50, 6, tzinfo=UTC))
    await app.state.services.processing.process_batch()

    response = await client.get("/timesheet", params={"month": "2025-09"})

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == 'attachment; filename="timesheet-2025-09.xlsx"'
    ws = sheet(response)
    assert [ws["A3"].value, ws["A4"].value, ws["A5"].value] == ["Антонов Иван", "Мехоношин Алексей", "ИТОГ:"]
    assert ws["T4"].value == 1  # колонка T — 19-е число
    assert ws["AG4"].value == "=SUM(B4:AF4)"
    assert ws["T5"].value == "=SUM(T3:T4)"


async def test_group_filter(app, client):
    for chat_id, title in [(CHAT_ID, "Пост 1"), (OTHER_CHAT_ID, "Пост 2")]:
        await client.post("/groups", json={"chat_id": chat_id, "title": title,
                                           "window_start": "07:00", "window_end": "10:00"})
    await client.post("/employees", json={"full_name": "Мехоношин Алексей"})
    await client.post("/employees", json={"full_name": "Петров Сергей"})
    await photo(app, CHAT_ID, 1, "Мехоношин", datetime(2025, 9, 19, 5, 0, tzinfo=UTC))
    await photo(app, OTHER_CHAT_ID, 1, "Петров", datetime(2025, 9, 20, 5, 0, tzinfo=UTC))
    await app.state.services.processing.process_batch()

    response = await client.get("/timesheet", params={"month": "2025-09", "chat_id": OTHER_CHAT_ID})

    ws = sheet(response)
    assert ws["B1"].value == "Табель за сентябрь 2025 — Пост 2"
    assert (ws["A3"].value, ws["U3"].value, ws["A4"].value) == ("Петров Сергей", 1, "ИТОГ:")
    assert response.headers["content-disposition"].endswith(f'timesheet-2025-09-{OTHER_CHAT_ID}.xlsx"')


@pytest.mark.parametrize(
    ("params", "code"),
    [({}, 422), ({"month": "2025-9"}, 422), ({"month": "2025-13"}, 422),
     ({"month": "2025-09", "chat_id": -1}, 404)],
    ids=["no-month", "bad-format", "bad-month", "unknown-group"],
)
async def test_errors(client, params, code):
    assert (await client.get("/timesheet", params=params)).status_code == code


async def test_requires_token(client):
    client.headers.pop("X-Admin-Token")
    assert (await client.get("/timesheet", params={"month": "2025-09"})).status_code == 401
