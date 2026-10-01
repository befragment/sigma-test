"""Формат xlsx по образцу листа «Табель» из ТЗ. Без БД: на входе готовый Timesheet."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.domain.models import Timesheet, TimesheetRow
from app.services.timesheet_xlsx import build_xlsx, day_column

SEPTEMBER = Timesheet(
    2025, 9, "Табель за сентябрь 2025",
    (TimesheetRow("Мехоношин Алексей", frozenset({19})), TimesheetRow("Петров Сергей", frozenset({1, 30}))),
)


def render(sheet: Timesheet):
    return load_workbook(BytesIO(build_xlsx(sheet)))["Табель"]


def test_header():
    ws = render(SEPTEMBER)

    assert ws["A2"].value == "Ф.И.О."
    assert [ws.cell(row=2, column=c).value for c in range(2, 33)] == list(range(1, 32))
    assert ws["AG2"].value == "кол.\nсмен"
    assert ws["B1"].value == "Табель за сентябрь 2025"
    assert "B1:AG1" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["A2"].font.b and ws["AG2"].font.b
    assert ws["A2"].fill.fgColor.rgb == "00D8D8D8"


def test_mark_lands_in_day_column():
    ws = render(SEPTEMBER)

    assert day_column(19) == "T"
    assert ws["A3"].value == "Мехоношин Алексей"
    assert ws["T3"].value == 1
    assert [ws.cell(row=3, column=c).value for c in range(2, 33)].count(1) == 1
    assert (ws["B4"].value, ws["AE4"].value, ws["C4"].value) == (1, 1, None)


def test_formulas():
    ws = render(SEPTEMBER)

    assert ws["AG3"].value == "=SUM(B3:AF3)"
    assert ws["AG4"].value == "=SUM(B4:AF4)"
    assert ws["A5"].value == "ИТОГ:"
    assert ws["B5"].value == "=SUM(B3:B4)"
    assert ws["T5"].value == "=SUM(T3:T4)"
    assert ws["AG5"].value == "=SUM(AG3:AG4)"
    assert ws["A5"].font.b
    assert ws.max_row == 5


def test_arial_everywhere():
    ws = render(SEPTEMBER)
    fonts = {cell.font.name for row in ws.iter_rows() for cell in row if cell.value is not None}
    assert fonts == {"Arial"}


@pytest.mark.parametrize(
    ("year", "month", "grey_days"),
    [(2025, 2, [29, 30, 31]), (2024, 2, [30, 31]), (2025, 9, [31]), (2025, 10, [])],
)
def test_missing_days_are_grey(year, month, grey_days):
    ws = render(Timesheet(year, month, "t", (TimesheetRow("Петров", frozenset()),)))

    grey = [day for day in range(1, 32) if ws[f"{day_column(day)}3"].fill.fgColor.rgb == "00A6A6A6"]
    assert grey == grey_days
    for day in grey_days:  # серый на всю высоту: шапка, сотрудники, итог
        assert {ws[f"{day_column(day)}{r}"].fill.fgColor.rgb for r in (2, 3, 4)} == {"00A6A6A6"}


def test_empty_timesheet_has_static_totals():
    ws = render(Timesheet(2025, 9, "t", ()))

    assert ws["A3"].value == "ИТОГ:"
    assert {ws.cell(row=3, column=c).value for c in range(2, 34)} == {0}
